"""
Word (.docx): JPG shablon sahifasi + ustiga matn (floating text) yoki tayyor raster.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm, Pt
from PIL import Image

from apps.articles.antiplagiat_template_render import (
    ASSETS_DIR,
    format_source_lines,
    render_certificate,
    render_report_cover,
    render_report_inner_page,
)

# A4 @ 96dpi reference; JPG 300dpi — Word da to'liq sahifa rasm
A4_PORTRAIT = (Mm(210), Mm(297))
A4_LANDSCAPE = (Mm(297), Mm(210))


def _blank_docx_with_fullpage_image(image_path: Path, *, landscape: bool = False) -> Document:
    """Faqat JPG fon — Word da qo'lda to'ldirish uchun shablon."""
    doc = Document()
    section = doc.sections[0]
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = A4_LANDSCAPE
    else:
        section.page_width, section.page_height = A4_PORTRAIT
    for margin in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(section, margin, Mm(0))
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(str(image_path), width=section.page_width, height=section.page_height)
    return doc


def build_empty_template_docx(kind: str) -> bytes:
    """kind: certificate | report_cover | report_inner"""
    mapping = {
        'certificate': (ASSETS_DIR / 'sertifikat.jpg', True),
        'report_cover': (ASSETS_DIR / 'hisobot1.jpg', False),
        'report_inner': (ASSETS_DIR / 'hisobot2.jpg', False),
    }
    path, landscape = mapping[kind]
    doc = _blank_docx_with_fullpage_image(path, landscape=landscape)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_filled_certificate_docx(values: dict[str, str]) -> bytes:
    """Tayyor JPG ustiga yozilgan — Word ga bitta rasm (1:1 ko'rinish)."""
    img = render_certificate(values, verify_id=values.get('document_number', ''))
    tmp = BytesIO()
    img.convert('RGB').save(tmp, format='JPEG', quality=95)
    tmp.seek(0)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = A4_LANDSCAPE
    for margin in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(section, margin, Mm(0))
    p = doc.add_paragraph()
    p.add_run().add_picture(tmp, width=section.page_width, height=section.page_height)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_filled_report_docx(
    cover_values: dict[str, str],
    inner_sections: list[tuple[str, list[str]]],
) -> bytes:
    """
    1-sahifa hisobot1, keyingi har bir bo'lim hisobot2 fonida.
    inner_sections: [(sarlavha, [qatorlar]), ...]
    """
    doc = Document()
    first = True
    for img in (
        render_report_cover(cover_values, verify_id=cover_values.get('document_number', '')),
        *[render_report_inner_page(lines, title=title) for title, lines in inner_sections],
    ):
        section = doc.sections[0] if first else doc.add_section()
        first = False
        section.page_width, section.page_height = A4_PORTRAIT
        for margin in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
            setattr(section, margin, Mm(0))
        buf = BytesIO()
        img.convert('RGB').save(buf, format='JPEG', quality=92)
        buf.seek(0)
        p = doc.add_paragraph()
        p.add_run().add_picture(buf, width=section.page_width, height=section.page_height)

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _dot_date(value) -> str:
    """ISO sana/vaqt → 08.10.2026 (bo'sh yoki noto'g'ri bo'lsa — bugungi sana)."""
    from datetime import datetime

    from django.utils import timezone

    if value:
        try:
            dt = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
            if timezone.is_aware(dt):
                dt = timezone.localtime(dt)
            return dt.strftime('%d.%m.%Y')
        except ValueError:
            pass
    return timezone.localtime().strftime('%d.%m.%Y')


def values_from_check_result(result: dict[str, Any], *, file_name: str = '') -> tuple[dict, dict, list]:
    report = result.get('report') or {}
    cert_no = str(report.get('certificate_number') or report.get('document_number') or '')
    enabled = len(report.get('enabled_module_ids') or [])
    module_names = [str(m).strip() for m in (report.get('search_modules') or []) if str(m).strip()]
    cert = {
        'document_number': cert_no,
        'check_date': report.get('check_date') or _dot_date(report.get('check_completed_at')),
        'author': report.get('author_name') or 'Muallif',
        'work_type': report.get('document_type') or 'Ilmiy ish',
        'file_name': file_name or report.get('document_name') or 'hujjat.docx',
        'citations': f"{float(report.get('citation_percent') or 0):.1f}%",
        'self_citation': f"{float(report.get('self_citation_percent') or 0):.1f}%",
        'plagiarism': f"{float(result.get('plagiarism_percentage') or 0):.2f}%",
        'originality': f"{float(result.get('originality') or 0):.2f}%",
        # «Qidiruv tizimlari»: tekshirilgan bazalar nomi; bo'lmasa soni
        'search_modules': ', '.join(module_names) if module_names else f'{enabled} ta bazada tekshirilgan',
    }
    cover = {
        'document_number': cert_no,
        'upload_date': report.get('check_date') or cert['check_date'],
        'author': cert['author'],
        'workplace': report.get('author_workplace') or '',
        'position': report.get('author_position') or '',
    }
    sources = result.get('sources') or report.get('sources') or []
    summary = [
        f"O'zlashtirish: {cert['plagiarism']} | Originallik: {cert['originality']} | "
        f"Iqtibos: {cert['citations']} | O'z-o'ziga: {cert['self_citation']}",
        f"Belgilar: {report.get('character_count', 0)} | Gaplar: {report.get('sentence_count', 0)}",
        '',
        'Manbalar (havola bilan):',
    ]
    lines = summary + format_source_lines(sources, start=1, limit=22)
    inner = [('Hujjat tekshirish natijalari', lines)]
    if len(sources) > 22:
        inner.append(('Manbalar (davomi)', format_source_lines(sources, start=23, limit=24)))
    if len(sources) > 46:
        inner.append(('Manbalar (davomi 2)', format_source_lines(sources, start=47, limit=24)))

    return cert, cover, inner
