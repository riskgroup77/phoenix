"""
"Barcha ilmiy ishlar bo'yicha ma'lumotnoma" (va bitta ish bo'yicha ma'lumotnoma) — PDF va DOCX.

Dizayn: assets/hisobot_template (Phoenix nashriyoti shabloni). 1-bet — muqova (PIL bilan chiziladi:
fon + logo + gulchambar + QR + muallif ma'lumotlari), keyingi betlar — fon ustida jadval.
Muqova rasm bo'lgani uchun PDF va DOCX da bir xil ko'rinadi.
"""
from __future__ import annotations

import copy
import io
import re
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from django.conf import settings
from django.utils import timezone

from config.doc_i18n import Dictionary, normalize_lang

ASSETS = Path(settings.BASE_DIR) / 'assets' / 'hisobot_template'
FONTS = ASSETS / 'fonts'

BRAND = (0x28, 0x6C, 0x81)  # sarlavha / yorliqlar (shablondan)
ACCENT = (0x39, 0xA1, 0xCE)  # vertikal chiziq
BRAND_HEX = '286C81'

# Ma'lumotnomaga kiradigan holatlar: jurnalga qabul qilingan va undan keyingilar (kitob — bosmaga berilgandan keyin)
REPORT_STATUSES = (
    'Accepted', 'NashrgaYuborilgan', 'Published',
    'SentToPrint', 'Printing', 'Ready', 'Packaging', 'Shipping', 'Delivered',
)
# Qabul ma'lumotnomasi va bitta ish hisoboti: yuborilgan (to'lovdan keyingi) har qanday holat, rad etilgandan tashqari
ACCEPTANCE_EXCLUDED_STATUSES = ('Draft', 'Rejected')

# Hujjat interfeys tilida (uz / ru / en) shakllanadi — config/doc_i18n.py
T = Dictionary({
    'title_all': ("BARCHA ILMIY ISHLAR BO‘YICHA MA’LUMOTNOMA", 'СПРАВКА ПО ВСЕМ НАУЧНЫМ РАБОТАМ',
                  'REFERENCE ON ALL RESEARCH WORKS'),
    'title_one': ("ILMIY ISH BO‘YICHA MA’LUMOTNOMA", 'СПРАВКА О НАУЧНОЙ РАБОТЕ', 'RESEARCH WORK REFERENCE'),
    'date': ('Hujjat yaratilgan sana:', 'Дата документа:', 'Date of issue:'),
    'number': ('Hujjat raqami:', 'Номер документа:', 'Document No.:'),
    'author': ('Muallif familya, ismi, sharifi:', 'Ф.И.О. автора:', 'Author\'s full name:'),
    'workplace': ('Muallif ish yoki o‘qish joyi:', 'Место работы или учёбы:', 'Place of work or study:'),
    'position': ('Muallifning lavozimi:', 'Должность автора:', 'Author\'s position:'),
    'default_position': ('Muallif', 'Автор', 'Author'),
    'qr': ('HUJJATNI TEKSHIRISH\nUCHUN QR KODDAN\nFOYDALANING', 'ДЛЯ ПРОВЕРКИ\nДОКУМЕНТА\nСКАНИРУЙТЕ QR-КОД',
           'SCAN THE QR CODE\nTO VERIFY\nTHIS DOCUMENT'),
    'col_no': ('№', '№', 'No.'),
    'col_title': ('Ilmiy ish nomi', 'Название научной работы', 'Title of the work'),
    'col_type': ('Ish turi', 'Вид работы', 'Type'),
    'col_publish': ('Nashr nomi va shakli', 'Издание и форма публикации', 'Publication and form'),
    'col_link': ('Internet havolasi', 'Интернет-ссылка', 'Internet link'),
    'col_coauthors': ('Hammualliflar', 'Соавторы', 'Co-authors'),
    'Maqola': ('Maqola', 'Статья', 'Article'),
    'Tezis': ('Tezis', 'Тезисы', 'Abstract'),
    'Kitob': ('Kitob', 'Книга', 'Book'),
    'Jurnal': ('Jurnal', 'Журнал', 'Journal'),
    'Konferensiya': ('Konferensiya', 'Конференция', 'Conference'),
    'Xalqaro konferensiya': ('Xalqaro konferensiya', 'Международная конференция', 'International conference'),
    'issue': ('{n}-son', '№ {n}', 'No. {n}'),
    'pages': ('{p}-betlar', 'с. {p}', 'pp. {p}'),
    'book_publisher': ('Kitob: Phoenix nashriyoti', 'Книга: издательство Phoenix', 'Book: Phoenix Publishing'),
    'in_progress': ('Nashr jarayonida', 'В процессе публикации', 'In the publication process'),
    'empty': ("Ma'lumotnomaga kiradigan ilmiy ishlar hozircha yo'q", 'Научных работ для справки пока нет',
              'No research works to include yet'),
    'total': ('Jami: {n} ta ilmiy ish.', 'Всего научных работ: {n}.', 'Total research works: {n}.'),
    'total_docx': ('Jami: {n} ta ilmiy ish.  Hujjat raqami: {num}', 'Всего научных работ: {n}.  Номер документа: {num}',
                   'Total research works: {n}.  Document No.: {num}'),
    'page': ('{n}-bet', 'стр. {n}', 'page {n}'),
    'publisher': ('Phoenix nashriyoti', 'Издательство Phoenix', 'Phoenix Publishing'),
    # AuthorPublication.publication_type
    'journal_local': ('Mahalliy OAK jurnali', 'Местный журнал ВАК', 'Local HAC journal'),
    'journal_international': ('Xalqaro OAK jurnali', 'Международный журнал ВАК', 'International HAC journal'),
    'journal_local_general': ('Mahalliy jurnal', 'Местный журнал', 'Local journal'),
    'journal_international_general': ('Xalqaro jurnal', 'Международный журнал', 'International journal'),
    'conference_local': ('Mahalliy konferensiya', 'Местная конференция', 'Local conference'),
    'conference_international': ('Xalqaro konferensiya', 'Международная конференция', 'International conference'),
    'scopus_journal': ('Scopus jurnali', 'Журнал Scopus', 'Scopus journal'),
    'scopus_conference': ('Scopus konferensiyasi', 'Конференция Scopus', 'Scopus conference'),
})


def columns(lang: str) -> tuple[str, ...]:
    return tuple(T.t(k, lang) for k in ('col_no', 'col_title', 'col_type', 'col_publish', 'col_link', 'col_coauthors'))


COLUMNS = columns('uz')
TITLE_ALL = T.t('title_all', 'uz')
TITLE_ONE = T.t('title_one', 'uz')


@dataclass
class WorkRow:
    title: str
    work_type: str
    publish: str
    link: str
    co_authors: list[str] = field(default_factory=list)
    sort_date: object = None


@dataclass
class ReportMeta:
    title: str
    number: str
    date: str
    author: str
    workplace: str
    position: str
    verify_url: str
    lang: str = 'uz'


# ---------------------------------------------------------------- ma'lumot yig'ish

def is_book_title(title: str | None) -> bool:
    return (title or '').strip().upper().startswith('[KITOB]')


def _is_standalone_plagiarism(article) -> bool:
    from .antiplagiat_utils import looks_like_standalone_antiplagiat

    return looks_like_standalone_antiplagiat(article)


def work_type_label(category_name: str, is_book: bool = False) -> str:
    if is_book:
        return 'Kitob'
    return 'Tezis' if 'konferens' in (category_name or '').lower() else 'Maqola'


def _venue_kind(category_name: str) -> str:
    c = (category_name or '').lower()
    if 'konferens' in c:
        return 'Xalqaro konferensiya' if ('xalqaro' in c or 'scopus' in c) else 'Konferensiya'
    return 'Jurnal'


def _fmt_date(d) -> str:
    if not d:
        return ''
    if hasattr(d, 'tzinfo') and getattr(d, 'hour', None) is not None:
        d = timezone.localtime(d) if timezone.is_aware(d) else d
    return d.strftime('%d.%m.%Y')


def _doi_link(doi: str) -> str:
    doi = (doi or '').strip()
    if not doi:
        return ''
    return doi if doi.lower().startswith('http') else f'https://doi.org/{doi}'


def _person_name(u) -> str:
    parts = [u.last_name, u.first_name, getattr(u, 'patronymic', '')]
    return ' '.join(p.strip() for p in parts if p and p.strip()) or u.email


def _article_row(article, lang: str = 'uz') -> WorkRow:
    journal = article.journal
    cat = getattr(getattr(journal, 'category', None), 'name', '') or ''
    book = is_book_title(article.title)
    title = re.sub(r'^\s*\[KITOB\]\s*', '', article.title or '', flags=re.I) if book else (article.title or '')
    when = article.issue.publication_date if article.issue_id and article.issue else article.submission_date
    if book:
        publish = T.t('book_publisher', lang)
    else:
        publish = f'{T.t(_venue_kind(cat), lang)}: “{journal.name}”'
        if article.issue_id and article.issue:
            publish += ', ' + T.t('issue', lang, n=article.issue.issue_number)
    if when:
        publish += f', {_fmt_date(when)}'
    link = _doi_link(article.doi) or (article.publication_url or '').strip()
    if not link and article.status == 'Published':
        link = f"{settings.PUBLIC_SITE_URL}/#/public/article/{article.id}"
    co = [_person_name(u) for u in article.co_authors.all()]
    return WorkRow(title=title, work_type=T.t(work_type_label(cat, book), lang), publish=publish,
                   link=link or T.t('in_progress', lang), co_authors=co, sort_date=when)


def _publication_row(pub, lang: str = 'uz') -> WorkRow:
    label = T.t(pub.publication_type, lang) if pub.publication_type in T.entries else ''
    is_conf = 'conference' in pub.publication_type
    venue = (pub.conference.title if is_conf and pub.conference_id else None) or (
        pub.journal.name if pub.journal_id else '') or (pub.conference.title if pub.conference_id else '')
    publish = f'{label}: “{venue}”' if venue else label
    if pub.pages:
        publish += ', ' + T.t('pages', lang, p=pub.pages)
    publish += f', {_fmt_date(pub.publication_date)}'
    co = [c.strip() for c in (pub.co_authors or '').split(',') if c.strip()]
    return WorkRow(title=pub.title, work_type=T.t('Tezis' if is_conf else 'Maqola', lang), publish=publish,
                   link=_doi_link(pub.doi) or pub.file_url or '—', co_authors=co, sort_date=pub.publication_date)


def collect_rows(user, article=None, lang: str = 'uz') -> list[WorkRow]:
    from apps.journals.models import AuthorPublication

    from .models import Article

    qs = Article.objects.select_related('journal__category', 'issue').prefetch_related('co_authors')
    if article is not None:
        return [_article_row(qs.get(pk=article.pk), lang)]
    rows = [
        _article_row(a, lang)
        for a in qs.filter(author=user, status__in=REPORT_STATUSES)
        if not _is_standalone_plagiarism(a)
    ]
    pubs = AuthorPublication.objects.filter(author=user).select_related('journal', 'conference')
    rows += [_publication_row(p, lang) for p in pubs]

    def key(r):
        d = r.sort_date
        if d is None:
            return ''
        return (timezone.localtime(d) if hasattr(d, 'hour') and timezone.is_aware(d) else d).strftime('%Y%m%d')

    rows.sort(key=key, reverse=True)
    return rows


def build_meta(user, *, article=None, workplace: str = '', position: str = '', lang: str = 'uz') -> ReportMeta:
    lang = normalize_lang(lang)
    if article is not None:
        number = f"ISH-{str(article.pk).replace('-', '')[:8].upper()}"
    else:
        number = f"HSB-{str(user.pk).replace('-', '')[:8].upper()}"
    return ReportMeta(
        title=T.t('title_one' if article is not None else 'title_all', lang),
        number=number,
        date=timezone.localtime().strftime('%d.%m.%Y'),
        author=_person_name(user),
        workplace=(workplace or user.affiliation or '').strip() or '—',
        position=(position or '').strip() or T.t('default_position', lang),
        verify_url=f"{settings.PUBLIC_SITE_URL}/#/verify/{number}",
        lang=lang,
    )


# ---------------------------------------------------------------- muqova (PIL)

COVER_W, COVER_H = 1654, 2339  # A4, 200 dpi


@lru_cache(maxsize=None)
def _cmap(path: str) -> frozenset:
    from reportlab.pdfbase.ttfonts import TTFontFile

    return frozenset(TTFontFile(path).charToGlyph.keys())


def _font_for(text: str, size: int, bold: bool = True):
    from PIL import ImageFont

    lex = str(FONTS / ('Lexend-Bold.ttf' if bold else 'Lexend-Regular.ttf'))
    path = lex if all(ord(ch) in _cmap(lex) for ch in text if not ch.isspace()) else str(
        FONTS / ('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf'))
    return ImageFont.truetype(path, size)


def _wrap(draw, text: str, size: int, max_w: int, max_lines: int = 3) -> list[str]:
    font = _font_for(text, size)
    words, lines, cur = text.split(), [], ''
    for w in words:
        trial = f'{cur} {w}'.strip()
        if draw.textlength(trial, font=font) <= max_w or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1].rstrip('.,;') + '…'
    return lines


def _paste_fit(canvas, img, box):
    """img (RGBA) ni box (x0, y0, x1, y1) ichiga nisbatni saqlab, markazga joylash."""
    x0, y0, x1, y1 = box
    scale = min((x1 - x0) / img.width, (y1 - y0) / img.height)
    im = img.resize((int(img.width * scale), int(img.height * scale)))
    px, py = x0 + ((x1 - x0) - im.width) // 2, y0 + ((y1 - y0) - im.height) // 2
    canvas.paste(im, (px, py), im)
    return px, py, scale


def render_cover(meta: ReportMeta) -> bytes:
    import qrcode
    from PIL import Image, ImageDraw

    W, H = COVER_W, COVER_H
    page = Image.open(ASSETS / 'page_bg.jpg').convert('RGB').resize((W, H), Image.LANCZOS)
    d = ImageDraw.Draw(page)

    # Sana va raqam
    s = 40
    y = int(0.112 * H)
    for x, label, value in ((int(0.10 * W), T.t('date', meta.lang), meta.date),
                            (int(0.58 * W), T.t('number', meta.lang), meta.number)):
        f = _font_for(label, s)
        width = d.textlength(label, font=f) + 14 + d.textlength(value, font=_font_for(value, s, bold=False))
        x = min(x, int(0.93 * W - width))  # uzun yorliq (ru/en) o'ng chetdan chiqmasin
        d.text((x, y), label, font=f, fill=BRAND)
        d.text((x + d.textlength(label, font=f) + 14, y), value, font=_font_for(value, s, bold=False), fill=BRAND)

    # Sarlavha (ikki qatorga bo'linadi, kenglikka sig'diriladi)
    words = meta.title.split()
    best = None
    for i in range(1, len(words)):
        a, b = ' '.join(words[:i]), ' '.join(words[i:])
        diff = abs(len(a) - len(b))
        if best is None or diff < best[0]:
            best = (diff, [a, b])
    lines = best[1] if best else [meta.title]
    size = 104
    while size > 50 and max(d.textlength(ln, font=_font_for(ln, size)) for ln in lines) > 0.84 * W:
        size -= 2
    ty = int(0.215 * H)
    for ln in lines:
        f = _font_for(ln, size)
        d.text(((W - d.textlength(ln, font=f)) / 2, ty), ln, font=f, fill=BRAND)
        ty += int(size * 1.18)

    # Muallif bloki: chapda yorliqlar, vertikal chiziq, o'ngda qiymatlar
    s = 42
    lx, bar_x, vx = int(0.05 * W), int(0.488 * W), int(0.512 * W)
    vmax = int(0.95 * W) - vx
    y = int(0.40 * H)
    top = y - 30
    for label, value in ((T.t('author', meta.lang), meta.author),
                         (T.t('workplace', meta.lang), meta.workplace),
                         (T.t('position', meta.lang), meta.position)):
        d.text((lx, y), label, font=_font_for(label, s), fill=BRAND)
        vlines = _wrap(d, value, s - 2, vmax)
        for i, ln in enumerate(vlines):
            d.text((vx, y + i * int(s * 1.25)), ln, font=_font_for(ln, s - 2, bold=False), fill=BRAND)
        y += max(1, len(vlines)) * int(s * 1.25) + 18
    d.rectangle((bar_x, top, bar_x + 7, max(y + 10, int(0.51 * H))), fill=ACCENT)

    # Gulchambar + QR (chapda), logo (o'ngda)
    # Gulchambar rasmi matnsiz (laurel_only.png) — "QR koddan foydalaning" yozuvi hujjat tilida chiziladi
    laurel = Image.open(ASSETS / 'laurel_only.png').convert('RGBA')
    lx0, ly0, sc = _paste_fit(page, laurel, (int(0.12 * W), int(0.70 * H), int(0.43 * W), int(0.80 * H)))
    lw, lh = int(laurel.width * sc), int(laurel.height * sc)
    qr = qrcode.QRCode(border=1, box_size=10, error_correction=qrcode.constants.ERROR_CORRECT_M)
    qr.add_data(meta.verify_url)
    qr.make(fit=True)
    qimg = qr.make_image(fill_color='#1f4f5f', back_color='white').convert('RGB')
    qs = int(lh * 0.72)
    qimg = qimg.resize((qs, qs), Image.NEAREST)
    page.paste(qimg, (int(lx0 + lw / 2 - qs / 2), int(ly0 + lh * 0.42 - qs / 2)))
    ty = ly0 + lh + 28
    for ln in T.t('qr', meta.lang).split('\n'):
        f = _font_for(ln, 38)
        d.text((lx0 + (lw - d.textlength(ln, font=f)) / 2, ty), ln, font=f, fill=BRAND)
        ty += 46
    _paste_fit(page, Image.open(ASSETS / 'logo.png').convert('RGBA'),
               (int(0.55 * W), int(0.70 * H), int(0.84 * W), int(0.89 * H)))

    out = io.BytesIO()
    page.save(out, format='JPEG', quality=90, dpi=(200, 200))
    return out.getvalue()


# ---------------------------------------------------------------- PDF

def _pdf_fonts() -> tuple[str, str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    names = (('PhxLexend-Bold', 'Lexend-Bold.ttf'), ('PhxSans', 'DejaVuSans.ttf'), ('PhxSans-Bold', 'DejaVuSans-Bold.ttf'))
    for name, file in names:
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(FONTS / file)))
    return 'PhxLexend-Bold', 'PhxSans', 'PhxSans-Bold'


def build_pdf(meta: ReportMeta, rows: list[WorkRow]) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import (BaseDocTemplate, Frame, NextPageTemplate, PageBreak, PageTemplate,
                                    Paragraph, Spacer, Table, TableStyle)

    head_font, body_font, body_bold = _pdf_fonts()
    PW, PH = A4
    cover = ImageReader(io.BytesIO(render_cover(meta)))
    bg = ImageReader(str(ASSETS / 'page_bg.jpg'))
    brand = colors.Color(*(c / 255 for c in BRAND))

    def on_cover(c, _doc):
        c.drawImage(cover, 0, 0, PW, PH)

    def on_body(c, _doc):
        c.drawImage(bg, 0, 0, PW, PH)
        c.setFont(body_font, 7.5)
        c.setFillColor(brand)
        c.drawCentredString(PW / 2, 2.85 * cm, f"{meta.number}  ·  {T.t('page', meta.lang, n=c.getPageNumber())}")

    buf = io.BytesIO()
    doc = BaseDocTemplate(buf, pagesize=A4, title=meta.title.title(), author=T.t('publisher', meta.lang))
    body_frame = Frame(1.3 * cm, 3.1 * cm, PW - 2.6 * cm, PH - 3.1 * cm - 3.3 * cm, id='body')
    doc.addPageTemplates([
        PageTemplate(id='cover', frames=[Frame(0, 0, PW, PH, id='c')], onPage=on_cover),
        PageTemplate(id='body', frames=[body_frame], onPage=on_body),
    ])

    h = ParagraphStyle('h', fontName=head_font, fontSize=16, leading=20, alignment=1, textColor=brand, spaceAfter=12)
    th = ParagraphStyle('th', fontName=body_bold, fontSize=8.5, leading=10.5, alignment=1, textColor=brand)
    td = ParagraphStyle('td', fontName=body_font, fontSize=8.5, leading=10.5, textColor=colors.HexColor('#1d3b47'))
    tdc = ParagraphStyle('tdc', parent=td, alignment=1)
    link = ParagraphStyle('lnk', parent=td, fontSize=7.5, leading=9.5, wordWrap='CJK')

    def p(text, style=td):
        return Paragraph(escape(text or '—').replace('\n', '<br/>'), style)

    data = [[p(c, th) for c in columns(meta.lang)]]
    for i, r in enumerate(rows, 1):
        co = '\n'.join(f'{k}. {n}' for k, n in enumerate(r.co_authors, 1)) if r.co_authors else '—'
        lk = r.link
        lk_par = (Paragraph(f'<link href="{escape(lk)}" color="#1f6f9a">{escape(lk)}</link>', link)
                  if lk.startswith('http') else p(lk, link))
        data.append([p(f'{i}.', tdc), p(r.title), p(r.work_type, tdc), p(r.publish), lk_par, p(co)])
    if len(data) == 1:
        data.append([p('—', tdc), p(T.t('empty', meta.lang)), '', '', '', ''])

    widths = [x * cm for x in (1.0, 4.2, 1.7, 4.0, 3.3, 3.3)]
    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ('GRID', (0, 0), (-1, -1), 0.6, brand),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e8f3f6')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f6fafb')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]
    if not rows:
        style.append(('SPAN', (1, 1), (-1, 1)))
    table.setStyle(TableStyle(style))

    total = ParagraphStyle('tot', fontName=body_font, fontSize=9, textColor=brand, spaceBefore=10)
    story = [NextPageTemplate('body'), PageBreak(), Paragraph(escape(meta.title), h), table,
             Spacer(1, 4), Paragraph(escape(T.t('total', meta.lang, n=len(rows))), total)]
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------- DOCX

def build_docx(meta: ReportMeta, rows: list[WorkRow]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Twips

    doc = Document(str(ASSETS / 'template.docx'))
    body = doc.element.body
    W_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    WP = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
    A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'

    # 1-bet: muqova rasmini almashtirish va butun sahifaga cho'zish
    anchor = body.find(f'.//{WP}anchor')
    blip = anchor.find(f'.//{A}blip')
    rid = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    doc.part.related_parts[rid]._blob = render_cover(meta)
    sect = doc.sections[0]
    cx, cy = int(sect.page_width), int(sect.page_height)
    for tag, rel in (('positionH', 'page'), ('positionV', 'page')):
        el = anchor.find(f'{WP}{tag}')
        el.set('relativeFrom', rel)
        el.find(f'{WP}posOffset').text = '0'
    anchor.find(f'{WP}extent').set('cx', str(cx))
    anchor.find(f'{WP}extent').set('cy', str(cy))
    anchor.set('behindDoc', '1')
    for ext in anchor.iter(f'{A}ext'):
        if ext.get('cx'):
            ext.set('cx', str(cx))
            ext.set('cy', str(cy))

    # 2-bet: sarlavha
    title_p = doc.paragraphs[1]
    for r in title_p.runs[1:]:
        r._r.getparent().remove(r._r)
    title_p.runs[0].text = meta.title
    title_p.runs[0].font.size = Pt(18)

    # Jadval: shablondagisini olib tashlab, 6 ustunli yangisini o'sha joyga qo'yamiz
    old = body.find(f'{W_NS}tbl')
    sect.left_margin, sect.right_margin = Twips(950), Twips(900)
    widths = (650, 2550, 1150, 2350, 1650, 1700)
    color = RGBColor.from_string(BRAND_HEX)
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    # Shablondagi "Table Grid" uslubi (id: a9)
    style_id = old.find(f'{W_NS}tblPr/{W_NS}tblStyle').get(qn('w:val'))
    table.style = next((s for s in doc.styles if s.style_id == style_id), None)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    def fill(cell, text, *, bold=False, center=False, size=10.5, width=None):
        cell.width = Twips(width)
        cell.paragraphs[0].text = ''
        for i, line in enumerate((text or '—').split('\n')):
            par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.space_after = Pt(0)
            run = par.add_run(line)
            run.font.name = 'Bookman Old Style'
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color

    for c, (name, w) in enumerate(zip(columns(meta.lang), widths)):
        fill(table.rows[0].cells[c], name, bold=True, center=True, size=11, width=w)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement('w:tblHeader')
    hdr.set(qn('w:val'), 'true')
    tr_pr.append(hdr)
    for i, r in enumerate(rows, 1):
        cells = table.add_row().cells
        co = '\n'.join(f'{k}. {n}' for k, n in enumerate(r.co_authors, 1)) if r.co_authors else '-'
        vals = (f'{i}.', r.title, r.work_type, r.publish, r.link, co)
        for c, (v, w) in enumerate(zip(vals, widths)):
            fill(cells[c], v, center=c in (0, 2), size=9 if c == 4 else 10.5, width=w)
    if not rows:
        cells = table.add_row().cells
        merged = cells[1].merge(cells[-1])
        fill(cells[0], '-', center=True, width=widths[0])
        fill(merged, T.t('empty', meta.lang), width=sum(widths[1:]))

    old.addprevious(table._tbl)
    old.getparent().remove(old)
    tbl_w = table._tbl.tblPr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        table._tbl.tblPr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')

    total = copy.deepcopy(title_p._p)
    table._tbl.addnext(total)
    from docx.text.paragraph import Paragraph as DocxParagraph

    tp = DocxParagraph(total, title_p._parent)
    tp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tp.runs[0].text = T.t('total_docx', meta.lang, n=len(rows), num=meta.number)
    tp.runs[0].font.size = Pt(10)
    tp.runs[0].font.bold = False

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
