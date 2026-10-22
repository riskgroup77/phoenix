"""
AI tarjima: «Oddiy» va «O'rta» turdagi buyurtmalar to'lovdan keyin avtomatik bajariladi.

  oddiy (basic)    — bitta o'tish: aniq, to'liq tarjima.
  o'rta (standard) — ikki o'tish: ilmiy uslubdagi tarjima, so'ng manba bilan solishtirib tahrir
                     (atamalar izchilligi, tushib qolgan joylar, ravonlik).
  puxta (expert)   — AI emas, tarjimon qo'lda bajaradi.

Hujjat paragraflarga ajratiladi va bo'laklab tarjima qilinadi (model javobi uzunligi cheklangan); natija DOCX.
AI ishlamasa (kalit yo'q, xato, bo'sh javob) — sababi ai_error ga yoziladi va buyurtma tarjimonlar navbatiga o'tadi:
muallif to'lagan xizmatsiz qolmaydi.
"""
from __future__ import annotations

import io
import logging
import os
import tempfile

from django.core.files.base import ContentFile
from django.db import close_old_connections
from django.utils import timezone

logger = logging.getLogger(__name__)

# Model uchun til nomlari (prompt inglizcha — barcha provayderlarda barqaror)
LANGUAGE_NAMES = {
    'uz': 'Uzbek (Latin script)',
    'en': 'English',
    'ru': 'Russian',
    'fr': 'French',
    'de': 'German',
    'es': 'Spanish',
    'ar': 'Arabic',
}
SEPARATOR = '@@@'
CHUNK_CHARS = 3500


class TranslationFailed(Exception):
    """Foydalanuvchiga/tarjimonga ko'rsatiladigan sabab bilan."""


def _lang(code: str) -> str:
    return LANGUAGE_NAMES.get((code or '').lower(), code or 'the target language')


def read_paragraphs(path: str) -> list[str]:
    """Hujjat matni paragraflar ro'yxati sifatida (bo'sh qatorlarsiz)."""
    if path.lower().endswith('.docx'):
        try:
            import docx

            document = docx.Document(path)
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        paragraphs.append(' | '.join(cells))
            if paragraphs:
                return paragraphs
        except Exception as exc:
            logger.warning('DOCX paragraflari o\'qilmadi, umumiy usul: %s', exc)
    from apps.services import extract_plain_text_from_file

    text = extract_plain_text_from_file(path) or ''
    return [line.strip() for line in text.splitlines() if line.strip()]


def chunk_paragraphs(paragraphs: list[str], limit: int = CHUNK_CHARS) -> list[list[str]]:
    chunks: list[list[str]] = []
    current: list[str] = []
    size = 0
    for paragraph in paragraphs:
        if current and size + len(paragraph) > limit:
            chunks.append(current)
            current, size = [], 0
        current.append(paragraph)
        size += len(paragraph)
    if current:
        chunks.append(current)
    return chunks


def _join(paragraphs: list[str]) -> str:
    return f'\n{SEPARATOR}\n'.join(paragraphs)


def _split(text: str, expected: int) -> list[str]:
    parts = [p.strip() for p in text.split(SEPARATOR)]
    parts = [p for p in parts if p]
    if len(parts) == expected:
        return parts
    # Model ajratgichni buzsa — matn yo'qolmasin: qatorlar bo'yicha
    return [line.strip() for line in text.splitlines() if line.strip() and line.strip() != SEPARATOR]


def _rules(source: str, target: str) -> str:
    return (
        f'Translate from {_lang(source)} to {_lang(target)}.\n'
        f'Paragraphs are separated by a line containing only {SEPARATOR}. Keep exactly the same number of '
        f'paragraphs and the same {SEPARATOR} separators, in the same order.\n'
        'Translate everything; do not summarize, skip, explain or add notes. Keep numbers, formulas, citations, '
        'references, URLs, proper names and abbreviations as appropriate for the target language.\n'
        'Return only the translated text.'
    )


def basic_prompt(text: str, source: str, target: str) -> str:
    return f'You are a professional translator.\n{_rules(source, target)}\n\nText:\n{text}'


def standard_prompt(text: str, source: str, target: str) -> str:
    return (
        'You are a professional translator of scientific and academic texts.\n'
        f'{_rules(source, target)}\n'
        'Use formal academic register and established scientific terminology of the target language; keep '
        'terminology consistent throughout; the result must read naturally to a native speaker.\n\n'
        f'Text:\n{text}'
    )


def review_prompt(source_text: str, draft: str, source: str, target: str) -> str:
    return (
        'You are a senior editor of scientific translations. Compare the draft translation with the source and '
        'produce the final version: fix mistranslations and omissions, unify terminology, improve fluency and '
        'academic style. Do not shorten the text.\n'
        f'{_rules(source, target)}\n\n'
        f'Source:\n{source_text}\n\nDraft translation:\n{draft}'
    )


def translate_paragraphs(service, paragraphs: list[str], source: str, target: str, quality: str) -> list[str]:
    from .models import TranslationRequest

    out: list[str] = []
    for chunk in chunk_paragraphs(paragraphs):
        text = _join(chunk)
        builder = standard_prompt if quality == TranslationRequest.QUALITY_STANDARD else basic_prompt
        draft = service.generate_text(builder(text, source, target))
        if not draft:
            raise TranslationFailed("AI javob bermadi (xizmat band yoki kalit yaroqsiz).")
        if quality == TranslationRequest.QUALITY_STANDARD:
            # Ikkinchi o'tish — tahrir; javob bo'lmasa birinchi variant qoladi
            draft = service.generate_text(review_prompt(text, draft, source, target)) or draft
        out.extend(_split(draft, len(chunk)))
    return out


def build_docx(paragraphs: list[str]) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _source_to_tempfile(request) -> str:
    ext = os.path.splitext(request.source_file_path.name or '')[1].lower()[:10] or '.docx'
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        request.source_file_path.open('rb')
        try:
            for chunk in request.source_file_path.chunks():
                tmp.write(chunk)
        finally:
            request.source_file_path.close()
        return tmp.name


def run_ai_translation(request_id) -> str:
    """Buyurtmani AI bilan tarjima qiladi. Qaytaradi: 'done' | 'failed' | 'skipped'."""
    from apps.services import get_gemini_service

    from .models import TranslationRequest
    from .notify import notify_ai_done, notify_ai_failed

    request = TranslationRequest.objects.select_related('author').filter(pk=request_id).first()
    if request is None or not request.is_ai or request.ai_error:
        return 'skipped'
    # Atomar "olish": to'lov signali ikki marta kelsa ham tarjima bir marta bajariladi
    if not TranslationRequest.objects.filter(pk=request.pk, status='Yangi').update(status='Jarayonda'):
        return 'skipped'
    request.status = 'Jarayonda'

    service = get_gemini_service()
    tmp_path = None
    try:
        if not service.available:
            raise TranslationFailed("AI xizmati sozlanmagan.")
        tmp_path = _source_to_tempfile(request)
        paragraphs = read_paragraphs(tmp_path)
        if not paragraphs:
            raise TranslationFailed("Hujjatdan matn ajratib bo'lmadi.")
        translated = translate_paragraphs(
            service, paragraphs, request.source_language, request.target_language, request.quality,
        )
        stem = os.path.splitext(os.path.basename(request.source_file_path.name or 'tarjima'))[0][:80]
        content = build_docx(translated)
        request.translated_file_path.save(f'{stem}_{request.target_language}.docx', ContentFile(content), save=False)
        request.status = 'Bajarildi'
        request.completion_date = timezone.now()
        request.ai_model = (service.active_model or '')[:100]
        request.save(update_fields=['translated_file_path', 'status', 'completion_date', 'ai_model'])
        notify_ai_done(request)
        return 'done'
    except Exception as exc:
        reason = str(exc) if isinstance(exc, TranslationFailed) else "AI tarjimada kutilmagan xato."
        if not isinstance(exc, TranslationFailed):
            logger.exception('AI tarjima xatosi %s', request_id)
        request.status = 'Yangi'
        request.ai_error = reason[:500]
        request.save(update_fields=['status', 'ai_error'])
        notify_ai_failed(request)
        return 'failed'
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.remove(tmp_path)


def start_ai_translation(request) -> None:
    """Fonda ishga tushirish: Celery worker bo'lsa — navbat, aks holda alohida oqim (thread)."""
    from django.conf import settings

    if not request.is_ai:
        return
    if getattr(settings, 'TRANSLATION_AI_SYNC', False):
        run_ai_translation(request.pk)
        return
    from .tasks import enqueue_ai_translation

    if enqueue_ai_translation(str(request.pk)):
        return
    import threading

    def _run():
        try:
            run_ai_translation(request.pk)
        finally:
            close_old_connections()

    threading.Thread(target=_run, name=f'ai-translate-{request.pk}', daemon=True).start()
