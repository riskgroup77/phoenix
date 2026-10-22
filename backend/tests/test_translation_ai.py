"""
Tarjima turlari: oddiy va o'rta — AI (to'lovdan keyin avtomatik), puxta — tarjimon (yangi «translator» roli).
AI provayderi soxta obyekt bilan almashtiriladi — haqiqiy API chaqirilmaydi.
"""
import io
import shutil
import tempfile
from decimal import Decimal
from unittest import mock

import docx
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from apps.notifications.models import Notification
from apps.payments.fulfillment import _fulfill_after_payment
from apps.payments.models import Transaction
from apps.translations.ai_translate import SEPARATOR
from apps.translations.models import TranslationRequest
from apps.users.models import User

MEDIA = tempfile.mkdtemp(prefix='phx-translation-')
PARAGRAPHS = ["Birinchi paragraf matni.", "Ikkinchi paragraf: natijalar va xulosalar.", "Uchinchi paragraf."]


def _source_docx(paragraphs=PARAGRAPHS):
    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    buf = io.BytesIO()
    document.save(buf)
    return SimpleUploadedFile('maqola.docx', buf.getvalue(),
                              'application/vnd.openxmlformats-officedocument.wordprocessingml.document')


class FakeAI:
    """Har bir paragrafni "EN: ..." ga aylantiradi; ajratgichlarni saqlaydi."""

    available = True
    active_model = 'fake-model'

    def __init__(self, fail=False):
        self.fail = fail
        self.prompts = []

    def generate_text(self, prompt):
        self.prompts.append(prompt)
        if self.fail:
            return None
        body = prompt.rsplit('Draft translation:\n', 1)[-1] if 'Draft translation:' in prompt else prompt.rsplit('Text:\n', 1)[-1]
        parts = [p.strip() for p in body.split(SEPARATOR)]
        return f'\n{SEPARATOR}\n'.join(p if p.startswith('EN: ') else f'EN: {p}' for p in parts)


def _user(phone, role):
    return User.objects.create_user(phone=phone, password='x-parol-123', email=f'{phone}@t.uz',
                                    first_name='Ism', last_name=role, affiliation='X', role=role)


@override_settings(MEDIA_ROOT=MEDIA)
class TranslationQualityTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.author = _user('998900008001', 'author')
        self.translator = _user('998900008002', 'translator')
        self.reviewer = _user('998900008003', 'reviewer')

    def api(self, user):
        c = APIClient()
        c.force_authenticate(user)
        return c

    def order(self, quality):
        res = self.api(self.author).post('/api/v1/translations/', {
            'title': 'Maqola', 'source_language': 'uz', 'target_language': 'en', 'quality': quality,
            'source_file_path': _source_docx(),
        }, format='multipart')
        self.assertEqual(res.status_code, 201, res.content)
        return TranslationRequest.objects.get(pk=res.data['id'])

    def pay(self, request):
        tx = Transaction.objects.create(user=self.author, translation_request=request, amount=request.cost,
                                        service_type='translation', status='completed')
        _fulfill_after_payment(tx)
        request.refresh_from_db()
        return request

    def translator_ids(self):
        data = self.api(self.translator).get('/api/v1/translations/').data
        data = data.get('results', data) if isinstance(data, dict) else data
        return {str(r['id']) for r in data}

    def test_analyze_returns_price_per_quality(self):
        res = self.api(self.author).post('/api/v1/translations/analyze_file/', {'file': _source_docx()},
                                         format='multipart')
        self.assertEqual(res.status_code, 200)
        q = {row['quality']: row for row in res.data['qualities']}
        self.assertEqual((q['basic']['price_per_word'], q['standard']['price_per_word'], q['expert']['price_per_word']),
                         (30, 60, 100))
        self.assertTrue(q['basic']['is_ai'] and q['standard']['is_ai'] and not q['expert']['is_ai'])
        words = res.data['word_count']
        self.assertEqual(q['standard']['cost'], words * 60)

    def test_basic_ai_translation_after_payment(self):
        fake = FakeAI()
        with mock.patch('apps.services.get_gemini_service', return_value=fake):
            request = self.order('basic')
            self.assertEqual(request.cost, Decimal(request.word_count * 30))
            self.assertEqual(request.status, 'Yangi')
            self.assertEqual(fake.prompts, [])  # to'lovgacha AI ishlamaydi
            request = self.pay(request)
        self.assertEqual(request.status, 'Bajarildi')
        self.assertEqual(request.ai_model, 'fake-model')
        self.assertEqual(len(fake.prompts), 1)  # bitta o'tish
        result = docx.Document(request.translated_file_path.path)
        self.assertEqual([p.text for p in result.paragraphs if p.text], [f'EN: {p}' for p in PARAGRAPHS])
        self.assertTrue(Notification.objects.filter(user=self.author, title='Tarjima tayyor').exists())
        self.assertNotIn(str(request.id), self.translator_ids())  # AI buyurtmasi tarjimon navbatida emas
        # To'lov signali qayta kelsa ham qayta tarjima qilinmaydi
        with mock.patch('apps.services.get_gemini_service', return_value=fake):
            _fulfill_after_payment(Transaction.objects.get(translation_request=request))
        self.assertEqual(len(fake.prompts), 1)

    def test_standard_uses_review_pass(self):
        fake = FakeAI()
        with mock.patch('apps.services.get_gemini_service', return_value=fake):
            request = self.pay(self.order('standard'))
        self.assertEqual(request.status, 'Bajarildi')
        self.assertEqual(len(fake.prompts), 2)  # tarjima + tahrir
        self.assertIn('senior editor', fake.prompts[1])
        self.assertEqual(request.cost, Decimal(request.word_count * 60))

    def test_ai_failure_goes_to_translator(self):
        with mock.patch('apps.services.get_gemini_service', return_value=FakeAI(fail=True)):
            request = self.pay(self.order('basic'))
        self.assertEqual(request.status, 'Yangi')
        self.assertTrue(request.ai_error)
        self.assertIn(str(request.id), self.translator_ids())
        self.assertTrue(Notification.objects.filter(user=self.translator, title='Yangi tarjima buyurtmasi').exists())
        self.assertTrue(Notification.objects.filter(user=self.author, title='Tarjima tarjimonga topshirildi').exists())
        # Tarjimon qo'lda bajaradi
        api = self.api(self.translator)
        res = api.patch(f'/api/v1/translations/{request.id}/', {'status': 'Jarayonda', 'reviewer': str(self.translator.id)},
                        format='json')
        self.assertEqual(res.status_code, 200, res.content)
        res = api.patch(f'/api/v1/translations/{request.id}/', {
            'status': 'Bajarildi', 'translated_file_path': _source_docx(['Translated.']),
        }, format='multipart')
        self.assertEqual(res.status_code, 200, res.content)

    def test_expert_order_for_translators_only_after_payment(self):
        with mock.patch('apps.services.get_gemini_service') as ai:
            request = self.order('expert')
            self.assertNotIn(str(request.id), self.translator_ids())  # to'lanmagan
            request = self.pay(request)
            ai.assert_not_called()
        self.assertEqual(request.status, 'Yangi')
        self.assertIn(str(request.id), self.translator_ids())
        self.assertTrue(Notification.objects.filter(user=self.translator, title='Yangi tarjima buyurtmasi').exists())
        # Taqrizchi endi tarjima navbatini ko'rmaydi
        data = self.api(self.reviewer).get('/api/v1/translations/').data
        data = data.get('results', data) if isinstance(data, dict) else data
        self.assertEqual(len(data), 0)

    def test_other_translator_cannot_take_assigned_order(self):
        with mock.patch('apps.services.get_gemini_service'):
            request = self.pay(self.order('expert'))
        request.reviewer = self.translator
        request.status = 'Jarayonda'
        request.save()
        other = _user('998900008004', 'translator')
        res = self.api(other).patch(f'/api/v1/translations/{request.id}/', {'status': 'Bajarildi'}, format='json')
        self.assertEqual(res.status_code, 403)

    def test_quality_cannot_change_after_order(self):
        with mock.patch('apps.services.get_gemini_service'):
            request = self.order('basic')
        admin = _user('998900008005', 'super_admin')
        self.api(admin).patch(f'/api/v1/translations/{request.id}/', {'quality': 'expert'}, format='json')
        request.refresh_from_db()
        self.assertEqual(request.quality, 'basic')

    def test_translator_cannot_submit_articles(self):
        res = self.api(self.translator).post('/api/v1/articles/', {'title': 'x'}, format='json')
        self.assertEqual(res.status_code, 403)
