"""
"Barcha ilmiy ishlar bo'yicha ma'lumotnoma" (PDF/DOCX), bitta ish ma'lumotnomasi, ish turi va QR tekshiruvi.
"""
import io
from decimal import Decimal

from django.core.cache import cache
from django.test import TestCase
from rest_framework.test import APIClient

from apps.articles import works_report
from apps.articles.models import Article
from apps.journals.models import Journal, JournalCategory
from apps.users.models import User

URL = '/api/v1/articles/reports/works/'


def _code8(pk) -> str:
    return str(pk).replace('-', '')[:8].upper()


class WorksReportTests(TestCase):
    def setUp(self):
        cache.clear()
        self.author = User.objects.create_user(phone='998900000501', password='x-parol-123', email='au@wr.uz',
                                               first_name='Ali', last_name='Valiyev', role='author',
                                               affiliation='TDPU')
        self.co = User.objects.create_user(phone='998900000502', password='x-parol-123', email='co@wr.uz',
                                           first_name='Bek', last_name='Alimov', role='author')
        ja = User.objects.create_user(phone='998900000503', password='x-parol-123', email='ja@wr.uz',
                                      first_name='J', last_name='A', role='journal_admin')
        conf = Journal.objects.create(name='MESAS', issn='1111-2222', description='d', journal_admin=ja,
                                      category=JournalCategory.objects.create(name='Xalqaro konferensiyalar'),
                                      publication_fee=Decimal('1'))
        jur = Journal.objects.create(name='Yangi tadqiqotlar', issn='2181-0325', description='d', journal_admin=ja,
                                     category=JournalCategory.objects.create(name='Mahalliy OAK jurnallari'),
                                     publication_fee=Decimal('1'))
        self.published = Article.objects.create(title='Конференция тезиси', abstract='a', journal=conf,
                                                author=self.author, status='Published')
        self.published.co_authors.add(self.co)
        self.new = Article.objects.create(title='Yangi maqola', abstract='a', journal=jur, author=self.author,
                                          status='Yangi')
        self.draft = Article.objects.create(title='Qoralama', abstract='a', journal=jur, author=self.author,
                                            status='Draft')
        self.client = APIClient()
        self.client.force_authenticate(self.author)

    def test_rows_only_accepted_and_later_with_work_type(self):
        rows = works_report.collect_rows(self.author)
        self.assertEqual([r.title for r in rows], ['Конференция тезиси'])
        self.assertEqual(rows[0].work_type, 'Tezis')
        self.assertEqual(rows[0].co_authors, ['Alimov Bek'])
        self.assertTrue(rows[0].publish.startswith('Xalqaro konferensiya: “MESAS”'))

    def test_pdf_and_docx_download(self):
        r = self.client.get(URL, {'fmt': 'pdf', 'position': 'Dotsent'})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r['Content-Type'], 'application/pdf')
        self.assertTrue(r.content.startswith(b'%PDF'))
        r = self.client.get(URL, {'fmt': 'docx'})
        self.assertEqual(r.status_code, 200)
        from docx import Document

        doc = Document(io.BytesIO(r.content))
        self.assertIn(works_report.TITLE_ALL, [p.text for p in doc.paragraphs])
        header = [c.text for c in doc.tables[0].rows[0].cells]
        self.assertEqual(tuple(header), works_report.COLUMNS)
        self.assertEqual(doc.tables[0].rows[1].cells[2].text, 'Tezis')

    def test_single_work_report_any_status_except_draft_rejected(self):
        r = self.client.get(URL, {'fmt': 'pdf', 'article': str(self.new.id)})
        self.assertEqual(r.status_code, 200)
        self.assertIn(f'ISH-{_code8(self.new.id)}', r['Content-Disposition'])
        self.assertEqual(self.client.get(URL, {'fmt': 'pdf', 'article': str(self.draft.id)}).status_code, 404)
        other = APIClient()
        other.force_authenticate(User.objects.create_user(phone='998900000504', password='x-parol-123',
                                                          email='x@wr.uz', first_name='X', last_name='Y'))
        self.assertEqual(other.get(URL, {'fmt': 'pdf', 'article': str(self.new.id)}).status_code, 404)

    def test_bad_format_and_anonymous(self):
        self.assertEqual(self.client.get(URL, {'fmt': 'xls'}).status_code, 400)
        self.assertIn(APIClient().get(URL).status_code, (401, 403))

    def test_verify_codes(self):
        guest = APIClient()
        r = guest.get(f'/api/v1/articles/verify/HSB-{_code8(self.author.pk)}/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['details']['works_count'], 1)
        r = guest.get(f'/api/v1/articles/verify/ISH-{_code8(self.new.pk)}/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data['details']['work_type'], 'Maqola')
        # Qabul ma'lumotnomasi: yangi yuborilgan maqolaga ham beriladi, qoralamaga — yo'q
        self.assertEqual(guest.get(f'/api/v1/articles/verify/QBL-{_code8(self.new.pk)}/').status_code, 200)
        self.assertEqual(guest.get(f'/api/v1/articles/verify/QBL-{_code8(self.draft.pk)}/').status_code, 404)

    def test_article_list_exposes_work_type(self):
        r = self.client.get('/api/v1/articles/')
        data = r.data.get('results', r.data) if isinstance(r.data, dict) else r.data
        types = {a['title']: a['work_type'] for a in data}
        self.assertEqual(types['Конференция тезиси'], 'Tezis')
        self.assertEqual(types['Yangi maqola'], 'Maqola')
